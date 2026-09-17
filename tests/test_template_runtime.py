from __future__ import annotations

import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches

ROOT = Path(__file__).resolve().parents[1]
SCRIPTS = ROOT / 'scripts'
sys.path.insert(0, str(SCRIPTS))

from fast_format_docx import replace_caption
from profile_config import infer_profile_from_template
from template_runtime import apply_table_rules, apply_template_sections, ensure_page_fields


def _pgnum(section, fmt: str):
    node = OxmlElement('w:pgNumType')
    node.set(qn('w:fmt'), fmt)
    node.set(qn('w:start'), '1')
    section._sectPr.append(node)


def _make_two_section_template(path: Path):
    doc = Document()
    front = doc.sections[0]
    front.left_margin = Inches(1.3)
    front.right_margin = Inches(0.8)
    front.header.paragraphs[0].text = 'FRONT HEADER'
    front.footer.paragraphs[0].text = 'FRONT FOOTER'
    _pgnum(front, 'upperRoman')

    custom_h1 = doc.styles.add_style('章标题', WD_STYLE_TYPE.PARAGRAPH)
    ppr = custom_h1._element.get_or_add_pPr()
    outline = OxmlElement('w:outlineLvl')
    outline.set(qn('w:val'), '0')
    ppr.append(outline)
    doc.add_paragraph('前置内容')

    body = doc.add_section(WD_SECTION.NEW_PAGE)
    body.left_margin = Inches(1.0)
    body.right_margin = Inches(1.0)
    body.header.is_linked_to_previous = False
    body.footer.is_linked_to_previous = False
    body.header.paragraphs[0].text = 'BODY HEADER'
    body.footer.paragraphs[0].text = 'BODY FOOTER'
    _pgnum(body, 'decimal')

    paragraph = doc.add_paragraph('章标题示例')
    paragraph.style = custom_h1
    custom_body = doc.styles.add_style('正文样式', WD_STYLE_TYPE.PARAGRAPH)
    paragraph = doc.add_paragraph('正文示例')
    paragraph.style = custom_body
    doc.save(path)


def _profile():
    return {
        'styles': {'heading_1': 'Heading 1', 'body': 'Body Text Generic'},
        'headings': {'max_level': 4},
        'pagination': {
            'front_matter': {'format': 'upperRoman'},
            'body': {'format': 'decimal'},
            'position': 'top_right',
            'page_field_required': True,
        },
        'captions': {
            'figure': {'label': '图', 'chapter_heading_level': 1, 'separator': '-'},
            'table': {'label': '表', 'chapter_heading_level': 1, 'separator': '-', 'repeat_header_on_continued_pages': False},
        },
    }


def test_custom_template_semantic_inference(tmp_path):
    template = tmp_path / 'template.docx'
    _make_two_section_template(template)
    inferred = infer_profile_from_template(_profile(), template)
    assert inferred['styles']['heading_1'] == '章标题'
    assert inferred['styles']['body'] == '正文样式'


def test_section_split_and_header_footer_transfer(tmp_path):
    template = tmp_path / 'template.docx'
    _make_two_section_template(template)
    src = Document()
    src.add_paragraph('目录')
    heading = src.add_paragraph('第一章')
    heading.style = 'Heading 1'
    src.add_paragraph('正文')
    source_path = tmp_path / 'source.docx'
    src.save(source_path)

    output = Document(source_path)
    result = apply_template_sections(output, template, _profile())
    assert result['section_split_added'] == 1
    assert len(output.sections) == 2
    assert output.sections[0].header.paragraphs[0].text == 'FRONT HEADER'
    assert output.sections[-1].header.paragraphs[0].text == 'BODY HEADER'
    assert output.sections[0].left_margin == Inches(1.3)
    assert output.sections[-1].left_margin == Inches(1.0)


def test_page_field_position_is_profile_driven(tmp_path):
    template = tmp_path / 'template.docx'
    _make_two_section_template(template)
    doc = Document()
    doc.add_paragraph('目录')
    heading = doc.add_paragraph('第一章')
    heading.style = 'Heading 1'
    apply_template_sections(doc, template, _profile())
    added = ensure_page_fields(doc, _profile())
    assert added >= 1
    assert any(
        'PAGE' in ''.join((node.text or '') for node in section.header._element.xpath('.//w:instrText'))
        for section in doc.sections
    )


def test_caption_alignment_is_inherited_from_template_style():
    doc = Document()
    caption_style = doc.styles.add_style('Left Caption', WD_STYLE_TYPE.PARAGRAPH)
    caption_style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph = doc.add_paragraph('图1-1 示例')
    paragraph.style = caption_style
    replace_caption(paragraph, {'label': '图', 'chapter': '1', 'sequence': '1', 'title': '示例'}, _profile())
    assert paragraph.alignment is None
    assert paragraph.style.paragraph_format.alignment == WD_ALIGN_PARAGRAPH.LEFT


def test_table_repeat_header_respects_profile_flag():
    doc = Document()
    table = doc.add_table(rows=2, cols=1)
    table.cell(0, 0).text = '表头'
    table.cell(1, 0).text = '内容'

    class Resolver:
        def resolve(self, _):
            return None

    apply_table_rules(doc, _profile(), Resolver())
    assert table.rows[0]._tr.get_or_add_trPr().find(qn('w:tblHeader')) is None


def test_source_custom_outline_style_is_recognized_as_heading():
    from fast_format_docx import classify
    doc = Document()
    custom = doc.styles.add_style('章节主标题', WD_STYLE_TYPE.PARAGRAPH)
    ppr = custom._element.get_or_add_pPr()
    outline = OxmlElement('w:outlineLvl')
    outline.set(qn('w:val'), '0')
    ppr.append(outline)
    paragraph = doc.add_paragraph('第一章 自定义标题')
    paragraph.style = custom
    role, _, _, state = classify(paragraph, 20, 'body', _profile())
    assert role == 'Heading 1'
    assert state == 'body'


def test_header_image_relationship_is_cloned(tmp_path):
    import base64
    image_path = tmp_path / 'dot.png'
    image_path.write_bytes(base64.b64decode(
        'iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAusB9Wl2nL8AAAAASUVORK5CYII='
    ))
    template = tmp_path / 'image-template.docx'
    donor = Document()
    run = donor.sections[0].header.paragraphs[0].add_run()
    run.add_picture(str(image_path), width=Inches(0.1))
    donor.sections[0].header.paragraphs[0].add_run(' LOGO')
    donor.add_paragraph('正文')
    donor.save(template)

    doc = Document()
    doc.add_paragraph('正文')
    apply_template_sections(doc, template, {'pagination': {'body': {'format': 'decimal'}}})
    output = tmp_path / 'out.docx'
    doc.save(output)
    reopened = Document(output)
    header = reopened.sections[0].header.paragraphs[0]
    assert header.text.strip() == 'LOGO'
    assert len(header._p.xpath('.//w:drawing')) == 1


def test_full_custom_template_pipeline_without_profile(tmp_path):
    from subprocess import run
    import sys
    template = tmp_path / 'custom.docx'
    doc = Document()
    for level, name in enumerate(['章标题', '节标题', '小节标题', '四级标题']):
        custom = doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        ppr = custom._element.get_or_add_pPr()
        outline = OxmlElement('w:outlineLvl')
        outline.set(qn('w:val'), str(level))
        ppr.append(outline)
    for name in ['正文样式', '文献条目', '特殊标题', '附录标题', '附录正文', '图题', '表题', '公式']:
        doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
    paragraph = doc.add_paragraph('章节示例'); paragraph.style = '章标题'
    paragraph = doc.add_paragraph('正文示例'); paragraph.style = '正文样式'
    paragraph = doc.add_paragraph('[1] 参考示例'); paragraph.style = '文献条目'
    doc.save(template)

    source = tmp_path / 'source.docx'
    doc = Document()
    doc.add_paragraph('1 绪论')
    doc.add_paragraph('正文引用[1]。')
    doc.add_paragraph('参考文献')
    doc.add_paragraph('[1] 张三. 示例[J]. 2026.')
    doc.save(source)
    output = tmp_path / 'output.docx'

    result = run([
        sys.executable, str(SCRIPTS / 'fast_format_docx.py'),
        '--source', str(source), '--output', str(output), '--template', str(template),
        '--strict-cross-references',
    ], cwd=ROOT, text=True, capture_output=True)
    assert result.returncode == 0, result.stderr
    formatted = Document(output)
    assert next(p.style.name for p in formatted.paragraphs if p.text == '绪论') == '章标题'
    assert next(p.style.name for p in formatted.paragraphs if p.text.startswith('张三')) == '文献条目'


def test_blank_template_header_clears_source_header(tmp_path):
    template = tmp_path / 'blank-header.docx'
    donor = Document()
    donor.add_paragraph('正文')
    donor.save(template)

    source = Document()
    source.sections[0].header.paragraphs[0].text = 'SHOULD DISAPPEAR'
    source.add_paragraph('正文')
    apply_template_sections(source, template, {'pagination': {'body': {'format': 'decimal'}}})
    assert all(not p.text.strip() for p in source.sections[0].header.paragraphs)
