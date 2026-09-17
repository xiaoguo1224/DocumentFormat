from __future__ import annotations

import json
import subprocess
import sys
import zipfile
from pathlib import Path

import yaml
from docx import Document
from lxml import etree

ROOT = Path(__file__).resolve().parents[1]
SCRIPTS = ROOT / 'scripts'
TEMPLATE = ROOT / 'templates/generic-formal/template.docx'
PROFILE = ROOT / 'templates/generic-formal/profile.yaml'
W = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'
NS = {'w': W}


def run(*args):
    return subprocess.run([sys.executable, *map(str, args)], cwd=ROOT, text=True, capture_output=True, check=True)


def make_source(path: Path):
    d = Document()
    d.add_paragraph('目录')
    d.add_paragraph('1 绪论')
    d.add_paragraph('已有研究表明该方法有效[1]，另一项工作见[2]。')
    d.add_paragraph('1.1 研究背景')
    d.add_paragraph('正文内容。')
    d.add_paragraph('参考文献')
    d.add_paragraph('[1] 张三. 示例文献一[J]. 测试学报, 2025.')
    d.add_paragraph('[2] 李四. 示例文献二[J]. 示例期刊, 2026.')
    d.add_paragraph('致谢')
    d.add_paragraph('感谢老师和同学。')
    d.add_paragraph('附录A')
    d.add_paragraph('附录正文。')
    d.save(path)


def format_file(src: Path, out: Path, profile=PROFILE, template=TEMPLATE):
    run(SCRIPTS/'fast_format_docx.py', '--source', src, '--output', out, '--template', template, '--profile', profile, '--ensure-toc', '--strict-cross-references')


def test_reference_numbering_and_body_ref(tmp_path):
    src=tmp_path/'in.docx';out=tmp_path/'out.docx';make_source(src);format_file(src,out)
    with zipfile.ZipFile(out) as z:
        doc=etree.fromstring(z.read('word/document.xml'));styles=etree.fromstring(z.read('word/styles.xml'));numbering=etree.fromstring(z.read('word/numbering.xml'))
    fields=doc.xpath('//w:instrText[contains(.,"REF _FmtRef")]/text()',namespaces=NS)
    assert len(fields)==2 and all('\\n' in x and '\\h' in x for x in fields)
    st=styles.xpath('//w:style[w:name/@w:val="References"]',namespaces=NS)[0]
    numid=st.xpath('./w:pPr/w:numPr/w:numId/@w:val',namespaces=NS)[0]
    aid=numbering.xpath(f'//w:num[@w:numId="{numid}"]/w:abstractNumId/@w:val',namespaces=NS)[0]
    lvl=numbering.xpath(f'//w:abstractNum[@w:abstractNumId="{aid}"]/w:lvl[@w:ilvl="0"]',namespaces=NS)[0]
    assert lvl.xpath('./w:lvlText/@w:val',namespaces=NS)==['[%1]']
    assert lvl.xpath('./w:suff/@w:val',namespaces=NS)==['tab']
    refs=doc.xpath('//w:p[w:pPr/w:pStyle/@w:val="References"]',namespaces=NS)
    assert [''.join(p.xpath('.//w:t/text()',namespaces=NS)) for p in refs]==['张三. 示例文献一[J]. 测试学报, 2025.','李四. 示例文献二[J]. 示例期刊, 2026.']


def test_section_state_exits_references(tmp_path):
    src=tmp_path/'in.docx';out=tmp_path/'out.docx';make_source(src);format_file(src,out)
    d=Document(out)
    pairs=[(p.text,p.style.name) for p in d.paragraphs]
    assert ('感谢老师和同学。','Body Text Generic') in pairs
    assert ('附录正文。','Appendix Body') in pairs


def test_repair_is_idempotent_for_reference_bookmarks(tmp_path):
    src=tmp_path/'in.docx';out=tmp_path/'out.docx';make_source(src);format_file(src,out)
    run(SCRIPTS/'repair_cross_references.py',out,out,'--profile',PROFILE,'--strict')
    with zipfile.ZipFile(out) as z:doc=etree.fromstring(z.read('word/document.xml'))
    names=doc.xpath('//w:bookmarkStart[starts-with(@w:name,"_FmtRef")]/@w:name',namespaces=NS)
    assert len(names)==2 and len(set(names))==2


def test_validator_reference_contract(tmp_path):
    src=tmp_path/'in.docx';out=tmp_path/'out.docx';make_source(src);format_file(src,out)
    result=run(SCRIPTS/'validate_docx.py',out,'--profile',PROFILE,'--require-toc','--require-multilevel','--require-crossrefs','--require-reference-numbering')
    data=json.loads(result.stdout)
    checks={x['check']:x['status'] for x in data['checks']}
    assert checks['reference native numbering']=='PASS'
    assert checks['reference typed prefixes']=='PASS'
    assert checks['bibliography body REF switch']=='PASS'


def test_profile_drives_custom_style_names(tmp_path):
    custom=tmp_path/'custom.docx'
    specs={
        'Heading1':('CustomH1','章标题'), 'Heading2':('CustomH2','节标题'),
        'Heading3':('CustomH3','小节标题'), 'Heading4':('CustomH4','四级标题'),
        'BodyTextGeneric':('CustomBody','正文样式'), 'References':('CustomRefs','文献条目'),
    }
    with zipfile.ZipFile(TEMPLATE) as zin, zipfile.ZipFile(custom,'w',zipfile.ZIP_DEFLATED) as zout:
        for info in zin.infolist():
            data=zin.read(info.filename)
            if info.filename=='word/styles.xml':
                root=etree.fromstring(data)
                for old_id,(new_id,new_name) in specs.items():
                    src=root.xpath(f'//w:style[@w:styleId="{old_id}"]',namespaces=NS)[0]
                    clone=etree.fromstring(etree.tostring(src))
                    clone.set(f'{{{W}}}styleId',new_id)
                    clone.xpath('./w:name',namespaces=NS)[0].set(f'{{{W}}}val',new_name)
                    root.append(clone)
                data=etree.tostring(root,xml_declaration=True,encoding='UTF-8',standalone='yes')
            zout.writestr(info,data)
    prof=yaml.safe_load(PROFILE.read_text(encoding='utf-8'))
    prof['styles']['heading_1']='章标题';prof['styles']['heading_2']='节标题';prof['styles']['heading_3']='小节标题';prof['styles']['heading_4']='四级标题';prof['styles']['body']='正文样式';prof['styles']['references']='文献条目'
    pp=tmp_path/'profile.yaml';pp.write_text(yaml.safe_dump(prof,allow_unicode=True,sort_keys=False),encoding='utf-8')
    src=tmp_path/'in.docx';out=tmp_path/'out.docx';make_source(src);format_file(src,out,pp,custom)
    d=Document(out)
    assert next(p.style.name for p in d.paragraphs if p.text=='绪论')=='章标题'
    assert next(p.style.name for p in d.paragraphs if p.text.startswith('张三'))=='文献条目'



def test_template_page_geometry_is_applied(tmp_path):
    src=tmp_path/'in.docx';out=tmp_path/'out.docx';make_source(src)
    source=Document(src)
    for sec in source.sections:
        sec.left_margin=1000000;sec.right_margin=1000000;sec.top_margin=1000000;sec.bottom_margin=1000000
    source.save(src)
    format_file(src,out)
    donor=Document(TEMPLATE).sections[-1]
    result=Document(out).sections[-1]
    assert result.left_margin==donor.left_margin
    assert result.right_margin==donor.right_margin
    assert result.top_margin==donor.top_margin
    assert result.bottom_margin==donor.bottom_margin
