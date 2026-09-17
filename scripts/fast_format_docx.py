#!/usr/bin/env python3
"""Fast, conservative formatter for the bundled generic formal DOCX template."""

from __future__ import annotations

import argparse
import json
import re
import shutil
import sys
import tempfile
import zipfile
from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph
from lxml import etree

from repair_cross_references import repair

W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
CT = "http://schemas.openxmlformats.org/package/2006/content-types"
DEFAULT_TEMPLATE = Path(__file__).resolve().parents[1] / "templates/generic-formal/template.docx"


def wt(name: str) -> str:
    return f"{{{W}}}{name}"


def merge_settings(source: bytes, template: bytes) -> bytes:
    root = etree.fromstring(source)
    donor = etree.fromstring(template)
    for name in ("captions", "updateFields"):
        old = root.find(wt(name))
        if old is not None:
            root.remove(old)
        new = donor.find(wt(name))
        if new is not None:
            root.append(deepcopy(new))
    update = root.find(wt("updateFields"))
    if update is None:
        update = etree.SubElement(root, wt("updateFields"))
    update.set(wt("val"), "true")
    return etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone="yes")


def apply_template_parts(template: Path, source: Path, output: Path) -> None:
    parts = {
        "word/styles.xml": "application/vnd.openxmlformats-officedocument.wordprocessingml.styles+xml",
        "word/numbering.xml": "application/vnd.openxmlformats-officedocument.wordprocessingml.numbering+xml",
        "word/theme/theme1.xml": "application/vnd.openxmlformats-officedocument.theme+xml",
        "word/fontTable.xml": "application/vnd.openxmlformats-officedocument.wordprocessingml.fontTable+xml",
    }
    with zipfile.ZipFile(template) as zt, zipfile.ZipFile(source) as zs:
        replacements = {name: zt.read(name) for name in parts if name in zt.namelist()}
        if "word/settings.xml" in zt.namelist() and "word/settings.xml" in zs.namelist():
            replacements["word/settings.xml"] = merge_settings(zs.read("word/settings.xml"), zt.read("word/settings.xml"))
        types = etree.fromstring(zs.read("[Content_Types].xml"))
        for name, content_type in parts.items():
            if name not in replacements:
                continue
            part_name = "/" + name
            node = next((n for n in types.findall(f"{{{CT}}}Override") if n.get("PartName") == part_name), None)
            if node is None:
                node = etree.SubElement(types, f"{{{CT}}}Override")
                node.set("PartName", part_name)
            node.set("ContentType", content_type)
        replacements["[Content_Types].xml"] = etree.tostring(types, xml_declaration=True, encoding="UTF-8", standalone="yes")
        with zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as zout:
            seen = set()
            for info in zs.infolist():
                seen.add(info.filename)
                zout.writestr(info, replacements.get(info.filename, zs.read(info.filename)))
            for name, data in replacements.items():
                if name not in seen:
                    zout.writestr(name, data)


def style_name(paragraph) -> str:
    try:
        return (paragraph.style.name or "").lower()
    except (KeyError, ValueError):
        return ""


def classify(paragraph, index: int, reference_mode: bool) -> tuple[str, str | None, dict | None, bool]:
    text = paragraph.text.strip()
    compact = re.sub(r"\s+", "", text)
    name = style_name(paragraph)
    if not text:
        return "Body Text Generic", None, None, reference_mode
    if compact in {"参考文献", "references"}:
        return "Special Heading 1", None, None, True
    if compact == "目录":
        return "TOC Title", None, None, reference_mode
    if compact in {"摘要", "中文摘要"}:
        return "Abstract Title CN", None, None, reference_mode
    if compact.lower() == "abstract":
        return "Abstract Title EN", None, None, reference_mode
    if compact in {"致谢", "结论"}:
        return "Special Heading 1", None, None, reference_mode
    if re.match(r"^附录[A-Z一二三四五六七八九十0-9]?", compact):
        return "Appendix Heading", None, None, reference_mode
    if reference_mode:
        return "References", None, None, True
    if "equation" in name or "公式" in name or paragraph._p.xpath(".//m:oMath | .//m:oMathPara"):
        return "Equation", None, None, reference_mode
    caption = re.match(r"^(图|表)\s*(\d+)(?:[-－.]([0-9]+))?\s*(.+)$", text)
    if caption:
        label, chapter, sequence, title = caption.groups()
        data = {"label": label, "chapter": chapter, "sequence": sequence or "1", "title": title.strip()}
        return ("Figure Caption" if label == "图" else "Table Caption"), None, data, reference_mode
    toc = re.search(r"toc\s*([1-3])", name)
    if toc:
        return f"TOC {toc.group(1)}", None, None, reference_mode
    heading = re.search(r"(?:heading|标题)\s*([1-4])", name)
    if heading:
        return f"Heading {heading.group(1)}", None, None, reference_mode
    if "title" in name and "subtitle" not in name:
        return "Title", None, None, reference_mode
    if "subtitle" in name:
        return "Subtitle", None, None, reference_mode
    if "list bullet" in name:
        return "List Bullet", None, None, reference_mode
    if "list number" in name:
        return "List Number", None, None, reference_mode
    typed = re.match(r"^(\d+(?:[.．]\d+){0,3})[ \t、.．]+(.+)$", text)
    if typed and len(text) <= 100 and not text.endswith(("。", "；", ";")):
        level = min(typed.group(1).replace("．", ".").count(".") + 1, 4)
        return f"Heading {level}", typed.group(2).strip(), None, reference_mode
    if re.match(r"^[•·●○▪-]\s*", text):
        return "List Bullet", re.sub(r"^[•·●○▪-]\s*", "", text), None, reference_mode
    numbered = re.match(r"^\d+[.)、]\s*(.+)$", text)
    if numbered:
        return "List Number", numbered.group(1).strip(), None, reference_mode
    if index < 12:
        sizes = [run.font.size.pt for run in paragraph.runs if run.font.size]
        if sizes and max(sizes) >= 18 and len(text) <= 80:
            return "Title", None, None, reference_mode
    return "Body Text Generic", None, None, reference_mode


def clear_direct_format(paragraph) -> None:
    ppr = paragraph._p.get_or_add_pPr()
    for name in ("jc", "spacing", "ind", "tabs", "pBdr", "shd", "numPr", "outlineLvl"):
        node = ppr.find(qn(f"w:{name}"))
        if node is not None:
            ppr.remove(node)
    for run in paragraph._p.xpath(".//w:r"):
        rpr = run.find(qn("w:rPr"))
        if rpr is None:
            continue
        for name in ("rFonts", "sz", "szCs", "color", "highlight", "spacing", "kern", "position"):
            node = rpr.find(qn(f"w:{name}"))
            if node is not None:
                rpr.remove(node)


def add_field(paragraph, instruction: str, cached: str) -> None:
    for kind, value in (("begin", None), (None, instruction), ("separate", None), (None, cached), ("end", None)):
        run = OxmlElement("w:r")
        if kind:
            node = OxmlElement("w:fldChar")
            node.set(qn("w:fldCharType"), kind)
        elif value == instruction:
            node = OxmlElement("w:instrText")
            node.set(qn("xml:space"), "preserve")
            node.text = value
        else:
            node = OxmlElement("w:t")
            node.text = value
        run.append(node)
        paragraph._p.append(run)


def replace_caption(paragraph, data: dict) -> None:
    for child in list(paragraph._p):
        if child.tag != qn("w:pPr"):
            paragraph._p.remove(child)
    paragraph.add_run(data["label"])
    add_field(paragraph, " STYLEREF 1 \\n ", data["chapter"])
    paragraph.add_run("-")
    add_field(paragraph, f" SEQ {data['label']} \\* ARABIC \\s 1 ", data["sequence"])
    paragraph.add_run(" " + data["title"])
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER


def has_field(root, field: str) -> bool:
    return any(field in (node.text or "").upper().split() for node in root.xpath(".//w:instrText"))


def ensure_page_field(doc: Document) -> int:
    if has_field(doc.element, "PAGE") or any(has_field(part.element, "PAGE") for part in doc.part.package.parts if hasattr(part, "element")):
        return 0
    touched = set()
    for section in doc.sections:
        footer = section.footer
        if id(footer.part) in touched:
            continue
        touched.add(id(footer.part))
        paragraph = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_field(paragraph, " PAGE ", "1")
    return len(touched)


def ensure_toc(doc: Document) -> bool:
    if has_field(doc.element, "TOC"):
        return False
    title = next((p for p in doc.paragraphs if style_name(p) == "toc title" or re.sub(r"\s+", "", p.text) == "目录"), None)
    if title is None:
        return False
    new_p = OxmlElement("w:p")
    title._p.addnext(new_p)
    paragraph = Paragraph(new_p, title._parent)
    paragraph.style = "TOC 1"
    add_field(paragraph, ' TOC \\o "1-3" \\h \\z \\u ', "打开文档后更新目录")
    return True


def format_docx(source: Path, template: Path, output: Path, add_toc: bool) -> dict:
    if source.resolve() == output.resolve():
        raise ValueError("output must differ from source")
    source_doc = Document(source)
    plans = []
    reference_mode = False
    for index, paragraph in enumerate(source_doc.paragraphs):
        role, replacement, caption, reference_mode = classify(paragraph, index, reference_mode)
        plans.append((role, replacement, caption))
    output.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix="docx-fast-format-") as temp:
        merged = Path(temp) / "merged.docx"
        apply_template_parts(template, source, merged)
        doc = Document(merged)
        if len(doc.paragraphs) != len(plans):
            raise RuntimeError("paragraph count changed while applying template")
        style_names = {style.name for style in doc.styles}
        stats = {"paragraphs": len(plans), "headings": 0, "captions": 0, "lists": 0, "fallback_styles": 0}
        for paragraph, (role, replacement, caption) in zip(doc.paragraphs, plans):
            target = role if role in style_names else "Normal"
            stats["fallback_styles"] += target == "Normal" and role != "Normal"
            paragraph.style = target
            clear_direct_format(paragraph)
            if replacement is not None:
                paragraph.text = replacement
            if caption:
                replace_caption(paragraph, caption)
                stats["captions"] += 1
            if role.startswith("Heading "):
                stats["headings"] += 1
            if role.startswith("List "):
                stats["lists"] += 1
        for table in doc.tables:
            if table.rows:
                tr_pr = table.rows[0]._tr.get_or_add_trPr()
                if tr_pr.find(qn("w:tblHeader")) is None:
                    tr_pr.append(OxmlElement("w:tblHeader"))
            for row in table.rows:
                for cell in row.cells:
                    for paragraph in cell.paragraphs:
                        if "Body Text Generic" in style_names:
                            paragraph.style = "Body Text Generic"
                        clear_direct_format(paragraph)
        stats["page_fields_added"] = ensure_page_field(doc)
        stats["toc_added"] = ensure_toc(doc) if add_toc else False
        doc.save(output)
    with zipfile.ZipFile(output) as zf:
        if zf.testzip() is not None:
            raise RuntimeError("generated DOCX failed ZIP integrity check")
    return stats


def main() -> None:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--template", type=Path, default=DEFAULT_TEMPLATE)
    parser.add_argument("--ensure-toc", action="store_true")
    parser.add_argument("--skip-cross-references", action="store_true")
    parser.add_argument("--strict-cross-references", action="store_true")
    args = parser.parse_args()
    for path in (args.source, args.template):
        if not path.is_file():
            parser.error(f"file not found: {path}")
    result = format_docx(args.source, args.template, args.output, args.ensure_toc)
    if not args.skip_cross_references:
        result["cross_references"] = repair(args.output, args.output, args.strict_cross_references)
    print(json.dumps(result, ensure_ascii=False))


if __name__ == "__main__":
    main()
