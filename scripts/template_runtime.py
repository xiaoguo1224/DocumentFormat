#!/usr/bin/env python3
from __future__ import annotations

from copy import deepcopy
from io import BytesIO
from pathlib import Path
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.opc.part import Part
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from lxml import etree

from profile_config import get, style

R = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'
REL_ATTRS = {f'{{{R}}}id', f'{{{R}}}embed', f'{{{R}}}link'}


def _copy_child(dst, src, tag: str) -> None:
    old = dst.find(qn(f'w:{tag}'))
    if old is not None:
        dst.remove(old)
    new = src.find(qn(f'w:{tag}'))
    if new is not None:
        dst.append(deepcopy(new))


def _field_exists(root, field: str) -> bool:
    needle = field.upper()
    return any(needle in (node.text or '').upper().split() for node in root.xpath('.//w:instrText'))


def _partname_template(partname) -> str:
    value = str(partname)
    directory, filename = value.rsplit('/', 1)
    m = re.match(r'^(.*?)(\d+)?(\.[^.]+)$', filename)
    if m:
        base, _, ext = m.groups()
        return f'{directory}/{base}%d{ext}'
    return f'{directory}/{filename}%d'


def _rewrite_rel_ids(blob: bytes, rid_map: dict[str, str]) -> bytes:
    if not rid_map:
        return blob
    try:
        root = etree.fromstring(blob)
    except etree.XMLSyntaxError:
        return blob
    changed = False
    for node in root.iter():
        for attr in list(node.attrib):
            if attr in REL_ATTRS and node.attrib[attr] in rid_map:
                node.attrib[attr] = rid_map[node.attrib[attr]]
                changed = True
    return etree.tostring(root, xml_declaration=True, encoding='UTF-8', standalone='yes') if changed else blob


def _clone_related_part(src_part, dst_package, cache: dict[int, Part]):
    key = id(src_part)
    if key in cache:
        return cache[key]
    if src_part.content_type.startswith('image/'):
        cloned = dst_package.get_or_add_image_part(BytesIO(src_part.blob))
        cache[key] = cloned
        return cloned

    new_name = dst_package.next_partname(_partname_template(src_part.partname))
    cloned = Part.load(new_name, src_part.content_type, src_part.blob, dst_package)
    cache[key] = cloned
    rid_map: dict[str, str] = {}
    for rid, rel in src_part.rels.items():
        if rel.is_external:
            new_rid = cloned.relate_to(rel.target_ref, rel.reltype, is_external=True)
        else:
            child = _clone_related_part(rel.target_part, dst_package, cache)
            new_rid = cloned.relate_to(child, rel.reltype)
        rid_map[rid] = new_rid
    cloned._blob = _rewrite_rel_ids(cloned.blob, rid_map)
    return cloned


def _clone_story(src_story, dst_story) -> None:
    """Clone one header/footer story including relationship-backed images/links."""
    dst_story.is_linked_to_previous = False
    src_part, dst_part = src_story.part, dst_story.part
    cache: dict[int, Part] = {}
    rid_map: dict[str, str] = {}
    for rid, rel in src_part.rels.items():
        if rel.is_external:
            new_rid = dst_part.relate_to(rel.target_ref, rel.reltype, is_external=True)
        else:
            child = _clone_related_part(rel.target_part, dst_part.package, cache)
            new_rid = dst_part.relate_to(child, rel.reltype)
        rid_map[rid] = new_rid

    cloned_root = etree.fromstring(etree.tostring(src_part.element))
    for node in cloned_root.iter():
        for attr in list(node.attrib):
            if attr in REL_ATTRS and node.attrib[attr] in rid_map:
                node.attrib[attr] = rid_map[node.attrib[attr]]
    dst_root = dst_part.element
    for child in list(dst_root):
        dst_root.remove(child)
    for child in list(cloned_root):
        dst_root.append(deepcopy(child))


def _section_pg_format(sect_pr) -> str | None:
    pg = sect_pr.find(qn('w:pgNumType'))
    return pg.get(qn('w:fmt')) if pg is not None else None


def _desired_format(profile, role: str) -> str | None:
    key = 'front_matter' if role == 'front' else 'body'
    return get(profile, 'pagination', key, 'format', default=None)


def _pick_donor_section(donor_sections, role: str, profile):
    desired = _desired_format(profile, role)
    if desired:
        for section in donor_sections:
            if _section_pg_format(section._sectPr) == desired:
                return section
    return donor_sections[0] if role == 'front' else donor_sections[-1]


def _paragraph_style_name(p) -> str:
    try:
        return p.style.name or ''
    except (KeyError, ValueError):
        return ''


def ensure_front_body_split(doc: Document, profile) -> int:
    """Create a front/body section split for one-section source documents when needed."""
    if len(doc.sections) != 1:
        return 0
    front_fmt = get(profile, 'pagination', 'front_matter', 'format', default=None)
    body_fmt = get(profile, 'pagination', 'body', 'format', default=None)
    if not front_fmt or not body_fmt or front_fmt == body_fmt:
        return 0
    h1 = style(profile, 'heading_1', 'Heading 1').casefold()
    paragraphs = doc.paragraphs
    index = next((i for i, p in enumerate(paragraphs) if _paragraph_style_name(p).casefold() == h1), None)
    if index is None or index <= 0:
        return 0
    previous = paragraphs[index - 1]
    ppr = previous._p.get_or_add_pPr()
    if ppr.find(qn('w:sectPr')) is not None:
        return 0
    ppr.append(deepcopy(doc.sections[-1]._sectPr))
    return 1


def _copy_document_setting(dst_doc: Document, src_doc: Document, tag: str) -> None:
    dst = dst_doc.settings._element
    src = src_doc.settings._element
    old = dst.find(qn(f'w:{tag}'))
    if old is not None:
        dst.remove(old)
    new = src.find(qn(f'w:{tag}'))
    if new is not None:
        dst.append(deepcopy(new))


def apply_template_sections(doc: Document, template: Path, profile) -> dict:
    donor = Document(template)
    donor_sections = list(donor.sections)
    if not donor_sections:
        return {'sections_updated': 0, 'stories_copied': 0, 'section_split_added': 0}

    split = ensure_front_body_split(doc, profile)
    _copy_document_setting(doc, donor, 'evenAndOddHeaders')
    _copy_document_setting(doc, donor, 'mirrorMargins')

    stories_copied = 0
    sections = list(doc.sections)
    for i, section in enumerate(sections):
        role = 'front' if i == 0 and len(sections) > 1 else 'body'
        src = _pick_donor_section(donor_sections, role, profile)
        dst_pr, src_pr = section._sectPr, src._sectPr
        for tag in ('pgSz', 'pgMar', 'cols', 'docGrid', 'pgNumType', 'titlePg'):
            _copy_child(dst_pr, src_pr, tag)

        title_page = src._sectPr.find(qn('w:titlePg')) is not None
        even_odd = donor.settings._element.find(qn('w:evenAndOddHeaders')) is not None
        story_pairs = (
            (src.header, section.header, True),
            (src.footer, section.footer, True),
            (src.first_page_header, section.first_page_header, title_page),
            (src.first_page_footer, section.first_page_footer, title_page),
            (src.even_page_header, section.even_page_header, even_odd),
            (src.even_page_footer, section.even_page_footer, even_odd),
        )
        for src_story, dst_story, enabled in story_pairs:
            has_content = bool(src_story.paragraphs and any(p.text or p._p.xpath('.//w:drawing|.//w:fldChar') for p in src_story.paragraphs))
            if enabled or has_content:
                # Clone even an explicitly blank default story so a blank template
                # header/footer clears source-document leftovers.
                _clone_story(src_story, dst_story)
                stories_copied += 1
    return {'sections_updated': len(sections), 'stories_copied': stories_copied, 'section_split_added': split}


def _add_simple_field(paragraph, instruction: str, cached: str = '1') -> None:
    for kind, value in (('begin', None), (None, instruction), ('separate', None), (None, cached), ('end', None)):
        run = OxmlElement('w:r')
        if kind:
            node = OxmlElement('w:fldChar'); node.set(qn('w:fldCharType'), kind)
        elif value == instruction:
            node = OxmlElement('w:instrText'); node.set(qn('xml:space'), 'preserve'); node.text = value
        else:
            node = OxmlElement('w:t'); node.text = value
        run.append(node); paragraph._p.append(run)


def ensure_page_fields(doc: Document, profile) -> int:
    if not bool(get(profile, 'pagination', 'page_field_required', default=True)):
        return 0
    position = str(get(profile, 'pagination', 'position', default='bottom_center')).casefold()
    use_header = position.startswith('top_')
    alignment = {
        'top_left': WD_ALIGN_PARAGRAPH.LEFT,
        'top_center': WD_ALIGN_PARAGRAPH.CENTER,
        'top_right': WD_ALIGN_PARAGRAPH.RIGHT,
        'bottom_left': WD_ALIGN_PARAGRAPH.LEFT,
        'bottom_center': WD_ALIGN_PARAGRAPH.CENTER,
        'bottom_right': WD_ALIGN_PARAGRAPH.RIGHT,
    }.get(position, WD_ALIGN_PARAGRAPH.CENTER)

    added = 0
    seen_parts: set[int] = set()
    for section in doc.sections:
        story = section.header if use_header else section.footer
        story.is_linked_to_previous = False
        if id(story.part) in seen_parts:
            continue
        seen_parts.add(id(story.part))
        if _field_exists(story.part.element, 'PAGE'):
            continue
        p = story.paragraphs[0] if story.paragraphs else story.add_paragraph()
        p.alignment = alignment
        _add_simple_field(p, ' PAGE ', '1')
        added += 1
    return added


def apply_table_rules(doc: Document, profile, resolver) -> dict:
    repeat_header = bool(get(profile, 'captions', 'table', 'repeat_header_on_continued_pages', default=False))
    table_style_name = get(profile, 'styles', 'table', default=None)
    cell_style_name = get(profile, 'styles', 'table_body', default=None)
    stats = {'tables': len(doc.tables), 'repeat_headers': 0, 'table_styles_applied': 0, 'cell_styles_applied': 0}
    for table in doc.tables:
        if table_style_name:
            try:
                table.style = table_style_name
                stats['table_styles_applied'] += 1
            except KeyError:
                pass
        if repeat_header and table.rows:
            trpr = table.rows[0]._tr.get_or_add_trPr()
            if trpr.find(qn('w:tblHeader')) is None:
                trpr.append(OxmlElement('w:tblHeader'))
            stats['repeat_headers'] += 1
        if cell_style_name:
            target = resolver.resolve(str(cell_style_name))
            if target is not None:
                for row in table.rows:
                    for cell in row.cells:
                        for p in cell.paragraphs:
                            p.style = target
                            stats['cell_styles_applied'] += 1
    return stats
